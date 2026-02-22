"""
Ипотечный калькулятор — веб-приложение на Flask.
Расчёт аннуитетного платежа и переплаты по кредиту.
Поддержка досрочных платежей: уменьшение платежа или срока.
"""

import io
from flask import Flask, render_template, request, jsonify, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)


def _annuity_payment(principal: float, months: int, r: float) -> float:
    """Аннуитетный платёж: P * (r*(1+r)^n) / ((1+r)^n - 1). При r=0: P/n."""
    if r == 0:
        return principal / months if months > 0 else 0.0
    rn = (1 + r) ** months
    return principal * (r * rn) / (rn - 1)


def mortgage_calculate(principal: float, years: int, rate_percent: float) -> dict:
    """
    Аннуитетный платёж (как в банковских калькуляторах):
    Коэффициент K = (r * (1+r)^n) / ((1+r)^n - 1), платёж = P * K.
    r — месячная ставка (годовая / 12 / 100).
    Платёж округляется до копеек, итог = платёж × количество месяцев.
    """
    if principal <= 0 or years <= 0:
        return {"error": "Сумма и срок должны быть больше нуля."}
    if rate_percent < 0:
        return {"error": "Ставка не может быть отрицательной."}

    months = years * 12
    # Месячная процентная ставка (номинальная, как в РФ)
    r = (rate_percent / 100) / 12

    if r == 0:
        monthly_exact = principal / months
    else:
        # Стандартная формула аннуитета
        r_plus_one_n = (1 + r) ** months
        annuity_coef = (r * r_plus_one_n) / (r_plus_one_n - 1)
        monthly_exact = principal * annuity_coef

    # Как в банках: платёж округляется до 2 знаков
    monthly_payment = round(monthly_exact, 2)
    if r == 0:
        # При 0% общая сумма выплат = сумма кредита, переплата = 0 (без погрешности округления)
        total_paid = principal
        overpayment = 0.0
    else:
        total_paid = monthly_payment * months
        overpayment = total_paid - principal

    # График погашения: помесячно платёж, основной долг, проценты, остаток
    schedule = []
    balance = principal
    if r == 0:
        principal_per_month = round(principal / months, 2)
        for month in range(1, months + 1):
            principal_part = principal_per_month if month < months else round(balance, 2)
            payment = principal_part
            interest = 0.0
            balance = round(balance - principal_part, 2)
            schedule.append({
                "month": month,
                "payment": round(payment, 2),
                "principal": round(principal_part, 2),
                "interest": 0.0,
                "balance": max(0, balance),
            })
    else:
        for month in range(1, months + 1):
            interest = balance * r
            if month == months:
                principal_part = balance
                payment = round(principal_part + interest, 2)
            else:
                principal_part = monthly_payment - interest
                payment = monthly_payment
            balance = balance - principal_part
            schedule.append({
                "month": month,
                "payment": round(payment, 2),
                "principal": round(principal_part, 2),
                "interest": round(interest, 2),
                "balance": round(max(0, balance), 2),
            })

    return {
        "monthly_payment": monthly_payment,
        "total_paid": round(total_paid, 2),
        "overpayment": round(overpayment, 2),
        "principal": round(principal, 2),
        "months": months,
        "schedule": schedule,
    }


def mortgage_calculate_with_early(
    principal: float,
    years: int,
    rate_percent: float,
    early_payments: list,
    reduction_type: str,
) -> dict:
    """
    Расчёт с досрочными платежами.
    early_payments: список { "month": int, "amount": float }.
    reduction_type: "payment" — уменьшать платёж, "term" — уменьшать срок.
    """
    if not early_payments:
        return mortgage_calculate(principal, years, rate_percent)

    months_total = years * 12
    r = (rate_percent / 100) / 12
    early_by_month = {}
    for e in early_payments:
        m = int(e.get("month", 0))
        a = float(e.get("amount", 0))
        if m > 0 and a > 0:
            early_by_month[m] = early_by_month.get(m, 0) + a

    if r == 0:
        monthly = round(principal / months_total, 2)
    else:
        monthly = round(_annuity_payment(principal, months_total, r), 2)

    schedule = []
    balance = principal
    total_paid_sum = 0.0
    month = 0
    max_months = months_total * 2

    while balance > 0 and month < max_months:
        month += 1
        if r == 0:
            principal_portion = min(monthly, balance)
            interest = 0.0
        else:
            interest = balance * r
            principal_portion = min(monthly - interest, balance)
            if principal_portion < 0:
                principal_portion = 0
            if balance - principal_portion < 0.02:
                principal_portion = balance

        payment = round(principal_portion + interest, 2)
        balance = round(balance - principal_portion, 2)
        early = min(early_by_month.get(month, 0), balance)
        if early > 0:
            balance = round(balance - early, 2)
            payment = round(payment + early, 2)
            if reduction_type == "payment" and balance > 0:
                remaining = months_total - month
                if remaining > 0:
                    if r == 0:
                        monthly = round(balance / remaining, 2)
                    else:
                        monthly = round(_annuity_payment(balance, remaining, r), 2)
        total_paid_sum += payment
        principal_total = principal_portion + early
        schedule.append({
            "month": month,
            "payment": payment,
            "principal": round(principal_total, 2),
            "interest": round(interest, 2),
            "balance": max(0, round(balance, 2)),
        })
        if balance <= 0:
            break

    overpayment = round(total_paid_sum - principal, 2)
    return {
        "monthly_payment": monthly,
        "total_paid": round(total_paid_sum, 2),
        "overpayment": overpayment,
        "principal": round(principal, 2),
        "months": len(schedule),
        "schedule": schedule,
    }


def mortgage_from_property_price(
    property_price: float,
    initial_payment: float,
    years: int,
    rate_percent: float,
    early_payments: list = None,
    reduction_type: str = "payment",
) -> dict:
    """Сумма кредита = стоимость недвижимости − первоначальный взнос. Расчёт с учётом взноса и досрочных платежей."""
    if property_price <= 0:
        return {"error": "Стоимость недвижимости должна быть больше нуля."}
    if initial_payment < 0:
        return {"error": "Первоначальный взнос не может быть отрицательным."}
    if initial_payment >= property_price:
        return {"error": "Первоначальный взнос не может быть больше или равен стоимости недвижимости."}
    principal = property_price - initial_payment
    early = early_payments or []
    if early:
        result = mortgage_calculate_with_early(principal, years, rate_percent, early, reduction_type)
    else:
        result = mortgage_calculate(principal, years, rate_percent)
    if "error" in result:
        return result
    result["property_price"] = round(property_price, 2)
    result["initial_payment"] = round(initial_payment, 2)
    result["total_cost"] = round(initial_payment + result["total_paid"], 2)
    result["down_payment_percent"] = round(initial_payment / property_price * 100, 1)
    result["required_monthly_income"] = round(result["monthly_payment"] / 0.4, 2)
    return result


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/calculate", methods=["POST"])
def calculate():
    data = request.get_json() or {}
    try:
        property_price = float(data.get("property_price", 0))
        initial_payment = float(data.get("initial_payment", 0))
        years = int(data.get("years", 0))
        rate = float(data.get("rate", 0))
        early_payments = data.get("early_payments") or []
        reduction_type = data.get("reduction_type", "payment")
        if reduction_type not in ("payment", "term"):
            reduction_type = "payment"
    except (TypeError, ValueError):
        return jsonify({"error": "Проверьте введённые данные."}), 400

    result = mortgage_from_property_price(
        property_price, initial_payment, years, rate,
        early_payments=early_payments,
        reduction_type=reduction_type,
    )
    if "error" in result:
        return jsonify(result), 400
    return jsonify(result)


@app.route("/export/excel", methods=["POST"])
def export_excel():
    """Принимает JSON с полем schedule (список строк графика), возвращает .xlsx файл."""
    data = request.get_json() or {}
    schedule = data.get("schedule") or []
    if not schedule:
        return jsonify({"error": "Нет данных для экспорта."}), 400
    wb = Workbook()
    ws = wb.active
    ws.title = "График платежей"
    headers = ["Месяц", "Платёж, ₽", "Основной долг, ₽", "Проценты, ₽", "Остаток долга, ₽"]
    thin = Side(style="thin", color="CCCCCC")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
    for row_idx, row in enumerate(schedule, 2):
        ws.cell(row=row_idx, column=1, value=row.get("month"))
        for col_idx, key in enumerate(["payment", "principal", "interest", "balance"], 2):
            val = row.get(key, 0)
            ws.cell(row=row_idx, column=col_idx, value=round(float(val), 2))
    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="grafik-platezhey.xlsx",
    )


if __name__ == "__main__":
    app.run(debug=True, port=5000)
